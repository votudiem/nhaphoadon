"""Local document extraction. No external AI or invoice uploads."""
import csv
from datetime import date, datetime
from decimal import Decimal
from html import escape
import io
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
import unicodedata
import zipfile
from openpyxl import load_workbook
from pypdf import PdfReader
from PIL import Image, ImageOps
import engine

ROOT=Path(__file__).parent
EXTENSIONS={'.xml','.pdf','.png','.jpg','.jpeg'}

def fold(value):
    return ''.join(c for c in unicodedata.normalize('NFD',str(value).lower().replace('đ','d')) if unicodedata.category(c)!='Mn')

def locale_number(value):
    text=str(value or '').strip().replace(' ','').replace('\u00a0','')
    text=re.sub(r'(?i)(vnd|usd|eur|₫|đ)$','',text)
    if not text:return ''
    if ',' in text:
        text=text.replace('.','').replace(',','.')
    elif re.fullmatch(r'-?\d{1,3}(?:\.\d{3})+',text):
        text=text.replace('.','')
    return str(engine.money(text))

ALIASES={
 'buyer_tax':['buyer_tax','mst nguoi mua','ma so thue nguoi mua'],
 'seller_tax':['seller_tax','mst nguoi ban','ma so thue nguoi ban','ma so thue'],
 'seller':['seller','ten nguoi ban','ten nguoi ban hang','ten don vi ban hang'],
 'serial':['serial','ky hieu hoa don','ky hieu'],
 'number':['number','so hoa don','so hd'],
 'date':['date','ngay hoa don','ngay lap','ngay lap hoa don'],
 'net':['net','tong tien chua thue','tong tien chua thue gtgt','tien truoc thue'],
 'vat':['vat','tong tien thue','tong tien thue gtgt','tien thue gtgt'],
 'total':['total','tong tien thanh toan','tong thanh toan','tong tien co thue'],
 'currency':['currency','don vi tien te','loai tien'],
 'rate':['rate','ty gia'],
 'kind':['kind','loai hoa don'],
 'status':['status','trang thai hoa don','trang thai'],
}

def tabular(raw,name):
    if Path(name).suffix.lower()=='.xlsx':
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if sum(i.file_size for i in z.infolist())>80_000_000:raise ValueError('Excel giải nén vượt 80 MB')
        wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
        try:
            ws=wb.active
            if ws.max_row and ws.max_row>50000:raise ValueError('Tối đa 50.000 dòng mỗi Excel')
            rows=[]
            for index,row in enumerate(ws.iter_rows()):
                if index>=50000:raise ValueError('Tối đa 50.000 dòng')
                values=[]
                for cell in row[:150]:
                    value=cell.value
                    if isinstance(value,(date,datetime)):value=value.strftime('%Y-%m-%d')
                    elif isinstance(value,(int,float)) and re.fullmatch('0+',cell.number_format or ''):
                        value=str(int(value)).zfill(len(cell.number_format))
                    values.append('' if value is None else str(value).strip())
                rows.append(values)
            return rows
        finally:wb.close()
    if Path(name).suffix.lower()!='.csv':raise ValueError('Danh sách CQT cần CSV hoặc XLSX; XLS cũ cần lưu lại thành XLSX')
    content=raw.decode('utf-8-sig')
    try:dialect=csv.Sniffer().sniff(content[:10000],delimiters=',;\t')
    except csv.Error:dialect=csv.excel
    return list(csv.reader(io.StringIO(content),dialect))

def preview_table(raw,name):
    try:rows=tabular(raw,name)
    except Exception as exc:raise ValueError('Không đọc được danh sách: '+str(exc)) from exc
    best=(-1,0,{})
    for index,row in enumerate(rows[:30]):
        mapping={key: next((i for i,v in enumerate(row) if fold(v).strip() in names),None) for key,names in ALIASES.items()}
        score=sum(value is not None for value in mapping.values())
        if score>best[0]:best=(score,index,mapping)
    if best[0]<3:raise ValueError('Không nhận diện được dòng tiêu đề. Dùng mẫu CSV hoặc sửa tên cột trong Excel.')
    _,index,mapping=best
    return {'headers':rows[index],'mapping':mapping,'header_row':index,'sample':rows[index+1:index+4],'row_count':sum(any(v.strip() for v in row) for row in rows[index+1:])}

def parse_table(raw,name,tax,period,mapping=None,header_row=None,confirm_buyer=False):
    preview=preview_table(raw,name)
    mapping=mapping or preview['mapping']; header_row=preview['header_row'] if header_row is None else int(header_row)
    rows=tabular(raw,name)
    if not 0<=header_row<min(30,len(rows)):raise ValueError('Dòng tiêu đề không hợp lệ')
    for field in ['seller_tax','serial','number','date','total']:
        if mapping.get(field) is None:raise ValueError(f'Chưa ánh xạ cột {field}')
    if mapping.get('buyer_tax') is None and not confirm_buyer:
        raise ValueError('File không có MST người mua. Cần xác nhận danh sách thuộc đúng công ty đang chọn.')
    records=[];filtered=0
    canonical=all(value in engine.FIELDS for value in rows[header_row])
    for line,row in enumerate(rows[header_row+1:],header_row+2):
        if not any(str(v).strip() for v in row):continue
        # CQT exports may end with a totals row. Only skip explicitly labeled summary rows.
        if fold(row[0]).strip() in ['tong cong','tong','total']:continue
        record={}
        for field in engine.FIELDS:
            col=mapping.get(field)
            if col is not None:
                if not isinstance(col,int) or not 0<=col<len(rows[header_row]):raise ValueError('Ánh xạ cột không hợp lệ')
                value=row[col] if col<len(row) else ''
            else:value=''
            record[field]=str(value).strip()
        if not record['buyer_tax'] and confirm_buyer:
            record['buyer_tax']=tax;record['buyer_asserted']=True
        if record['buyer_tax']!=tax:raise ValueError(f'Dòng {line}: MST người mua khác công ty đang chọn')
        date_value=record['date']
        if re.fullmatch(r'\d{1,2}/\d{1,2}/\d{4}',date_value):date_value=datetime.strptime(date_value,'%d/%m/%Y').strftime('%Y-%m-%d')
        record['date']=date_value[:10]
        try:
            if not canonical:
                for key in ['net','vat','total','rate']:record[key]=locale_number(record[key])
            record['currency']=record['currency'].upper() or 'VND'
            if 'ban hang' in fold(record['kind']):record['kind']='sales'
            if fold(record['status']) in ['hoa don goc','hoa don moi','goc','new','1']:record['status']=''
            engine.validate(record)
        except (ValueError,TypeError) as exc:raise ValueError(f'Dòng {line}: {exc}') from exc
        if record['date'][:7]!=period:filtered+=1;continue
        records.append(record)
    if not records:raise ValueError('Không có hóa đơn đúng kỳ đang chọn')
    keys=[engine.identity(r) for r in records]
    if len(keys)!=len(set(keys)):raise ValueError('Danh sách có khóa hóa đơn trùng. Kiểm tra trước khi nhập.')
    return records,filtered

def ocr_image(image):
    exe=shutil.which('tesseract')
    if not exe:raise ValueError('Chưa cài Tesseract OCR. Có thể nhập tay các trường sau khi xem file gốc.')
    with tempfile.TemporaryDirectory() as directory:
        path=Path(directory)/'page.png'
        image=ImageOps.exif_transpose(image).convert('RGB');image.thumbnail((2600,3600));image.save(path)
        langs=subprocess.run([exe,'--list-langs'],capture_output=True,text=True,timeout=15)
        lang='vie+eng' if 'vie' in langs.stdout.splitlines() else 'eng'
        command=[exe,str(path),'stdout','-l',lang,'--psm','3']
        result=subprocess.run(command,capture_output=True,text=True,encoding='utf-8',errors='replace',timeout=90)
        if result.returncode:raise ValueError('OCR không đọc được file: '+result.stderr[-250:])
        return result.stdout,lang

def extract_text(raw,name):
    ext=Path(name).suffix.lower()
    if ext=='.pdf':
        reader=PdfReader(io.BytesIO(raw),strict=False)
        if reader.is_encrypted and not reader.decrypt(''):raise ValueError('PDF có mật khẩu')
        if len(reader.pages)>50:raise ValueError('PDF vượt 50 trang; hãy tách theo hóa đơn')
        content='\n'.join(page.extract_text() or '' for page in reader.pages)
        if len(re.sub(r'\s','',content))>=60:return content[:200000],'pdf-text'
        import pypdfium2 as pdfium
        doc=pdfium.PdfDocument(raw)
        try:
            if len(doc)>5:raise ValueError('PDF scan vượt 5 trang; cần tách file trước khi OCR')
            texts=[];lang=''
            for index in range(len(doc)):
                page=doc[index]
                try:
                    bitmap=page.render(scale=2)
                    try:part,lang=ocr_image(bitmap.to_pil());texts.append(part)
                    finally:bitmap.close()
                finally:page.close()
            return '\n'.join(texts),f'ocr:{lang}'
        finally:doc.close()
    if ext in {'.png','.jpg','.jpeg'}:
        with Image.open(io.BytesIO(raw)) as image:
            if image.width*image.height>40_000_000:raise ValueError('Ảnh vượt 40 megapixel')
            result,lang=ocr_image(image)
            return result,f'ocr:{lang}'
    raise ValueError('Định dạng không được hỗ trợ')

def parse_text(content,method):
    """Conservative candidates only; every non-XML extraction requires human approval."""
    original=[x.strip() for x in content.splitlines() if x.strip()]
    lines=[fold(x) for x in original]
    result={field:'' for field in engine.FIELDS};result.update(currency='VND',kind='vat',items=[],method=method,needs_review=True,extracted_text=content)
    def find(pattern):
        for raw,line in zip(original,lines):
            match=re.search(pattern,line,re.I)
            if match:return raw[match.start(1):match.end(1)].strip()
        return ''
    result['number']=find(r'(?:so hoa don|invoice no\.?|so \(no\.?\)|so)\s*[:：]\s*(\d{1,20})')
    result['serial']=find(r'(?:ky hieu(?: hoa don)?|serial(?: no\.?)?)\s*[:：]\s*([a-z0-9]{4,15})')
    date_text=find(r'(?:ngay(?: lap)?|date)\s*[:：]?\s*(\d{1,2}[/.-]\d{1,2}[/.-]\d{4})')
    if date_text:
        result['date']=datetime.strptime(re.sub(r'[.-]','/',date_text),'%d/%m/%Y').strftime('%Y-%m-%d')
    else:
        match=re.search(r'ngay\s+(\d{1,2})\s+thang\s+(\d{1,2})\s+nam\s+(\d{4})',fold(content))
        if match:
            try:result['date']=date(int(match[3]),int(match[2]),int(match[1])).isoformat()
            except ValueError:pass
    tax_matches=[]
    for line in lines:
        match=re.search(r'(?:ma so thue|tax code|mst)\s*[:：]?\s*(\d{10}(?:-?\d{3})?)',line)
        if match:tax_matches.append(match[1])
    if tax_matches:result['seller_tax']=tax_matches[0]
    if len(tax_matches)>1:result['buyer_tax']=tax_matches[1]
    result['seller']=find(r'(?:don vi ban hang|ten nguoi ban|seller)\s*[:：]\s*(.+)')
    labels={'net':r'(?:tong tien chua thue(?: gtgt)?|cong tien hang|tien truoc thue|subtotal)\s*[:：]?\s*([\d., ]+)',
            'vat':r'(?:tong tien thue(?: gtgt)?|tien thue gtgt|vat amount)\s*[:：]?\s*([\d., ]+)',
            'total':r'(?:tong (?:cong )?tien thanh toan|tong thanh toan|grand total|total payment)\s*[:：]?\s*([\d., ]+)',
            'rate':r'(?:ty gia|exchange rate)\s*[:：]?\s*([\d., ]+)'}
    for field,pattern in labels.items():
        value=find(pattern)
        if value:
            try:result[field]=locale_number(value)
            except ValueError:pass
    currency=find(r'(?:don vi tien te|currency)\s*[:：]\s*([a-z]{3})')
    if currency:result['currency']=currency.upper()
    if 'hoa don ban hang' in fold(content):result['kind']='sales'
    # Item descriptions cannot be safely inferred from arbitrary PDF layout.
    return result

def parse_document(raw,name):
    if Path(name).suffix.lower()=='.xml':
        record=engine.parse_xml(raw);record.update(method='xml',needs_review=False);return record
    content,method=extract_text(raw,name)
    if 'ban the hien duoc tao tu dong tu du lieu xml' in fold(content):raise ValueError('DERIVED_PDF: PDF do ứng dụng tạo, không đưa lại vào engine')
    if not content.strip():raise ValueError('File không có văn bản đọc được')
    return parse_text(content,method)

def pdf_representation(invoice):
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    if 'InvoiceUnicode' not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont('InvoiceUnicode',str(ROOT/'assets'/'DejaVuSans.ttf')))
    normal=ParagraphStyle('normal',fontName='InvoiceUnicode',fontSize=9,leading=15)
    heading=ParagraphStyle('heading',parent=normal,fontSize=17,leading=24,spaceAfter=15)
    def p(value):return Paragraph(escape(str(value or '—')).replace('\n','<br/>'),normal)
    out=io.BytesIO();doc=SimpleDocTemplate(out,pagesize=A4,rightMargin=32,leftMargin=32,topMargin=32,bottomMargin=32)
    story=[Paragraph('HÓA ĐƠN · BẢN THỂ HIỆN',heading),p('Bản thể hiện được tạo tự động từ dữ liệu XML hóa đơn điện tử. Không phải PDF gốc của nhà cung cấp; chưa xác minh chữ ký số.'),Spacer(1,15)]
    labels={'kind':'Loại hóa đơn','serial':'Ký hiệu','number':'Số hóa đơn','date':'Ngày hóa đơn','seller':'Người bán','seller_tax':'MST người bán','buyer':'Người mua','buyer_tax':'MST người mua','currency':'Đơn vị tiền tệ','rate':'Tỷ giá'}
    for field,label in labels.items():story.append(p(f'{label}: {invoice.get(field,"")}'))
    story.append(Spacer(1,15))
    rows=[[p(x) for x in ['Hàng hóa / dịch vụ','SL','ĐVT','Đơn giá','Thành tiền']]]
    for item in invoice.get('items',[]):rows.append([p(item.get(k,'')) for k in ['name','quantity','unit','price','amount']])
    table=Table(rows,colWidths=[245,45,45,65,125],repeatRows=1,hAlign='LEFT')
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#edf3ec')),('GRID',(0,0),(-1,-1),.3,colors.HexColor('#d8e2d5')),('VALIGN',(0,0),(-1,-1),'TOP'),('BOTTOMPADDING',(0,0),(-1,-1),8)]))
    story.extend([table,Spacer(1,15)])
    for field,label in [('net','Tiền trước thuế'),('vat','VAT'),('total','Tổng thanh toán')]:story.append(p(f'{label}: {invoice.get(field,"")}'))
    doc.build(story);return out.getvalue()
