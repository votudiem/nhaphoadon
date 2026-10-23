import base64
import copy
import csv
import io
import json
from pathlib import Path
import tempfile
import threading
import unittest
from urllib.request import Request, urlopen
from urllib.error import HTTPError
from http.server import ThreadingHTTPServer
from openpyxl import load_workbook
import app
import engine

XML = '''<HDon><DLHDon><TTChung><KHMSHDon>1</KHMSHDon><KHHDon>1K26THT</KHHDon><SHDon>00000057</SHDon><NLap>2026-10-03</NLap><DVTTe>VND</DVTTe></TTChung><NDHDon><NBan><Ten>Công ty An Phát</Ten><MST>0311111111</MST></NBan><NMua><MST>0317557745</MST></NMua><DSHHDVu><HHDVu><THHDVu>Thiết bị giám sát hành trình GT-S7</THHDVu><SLuong>2</SLuong><DVTinh>Bộ</DVTinh><ThTien>1000000</ThTien></HHDVu></DSHHDVu><TToan><TgTCThue>1000000</TgTCThue><TgTThue>80000</TgTThue><TgTTTBSo>1080000</TgTTTBSo></TToan></NDHDon></DLHDon></HDon>'''.encode()
TAX='0317557745'
PERIOD='2026-10'

def csv_bytes(records):
    out=io.StringIO();writer=csv.DictWriter(out,fieldnames=engine.FIELDS);writer.writeheader()
    for record in records:writer.writerow({k:record.get(k,'') for k in engine.FIELDS})
    return out.getvalue().encode('utf-8-sig')

def upload(source,raw,name):
    return {'tax':TAX,'period':PERIOD,'source':source,'files':[{'name':name,'data':base64.b64encode(raw).decode()}]}

class EngineTests(unittest.TestCase):
    def setUp(self):self.doc=engine.parse_xml(XML)
    def test_string_identity_and_series(self):
        source=dict(self.doc,serial='K26THT',number='57')
        self.assertEqual(engine.identity(source),engine.identity(self.doc))
        result,_=engine.reconcile([source],[self.doc]);self.assertEqual(result[0]['issues'],[])
        self.assertEqual(self.doc['number'],'00000057')
    def test_ambiguous_missing_orphan_and_wrong_buyer(self):
        other=dict(self.doc,buyer_tax='9999999999')
        result,orphans=engine.reconcile([self.doc],[other]);self.assertEqual(result[0]['issues'],['CHƯA TÌM THẤY CHỨNG TỪ']);self.assertEqual(orphans,[other])
        result,_=engine.reconcile([self.doc],[self.doc,copy.deepcopy(self.doc)])
        self.assertEqual(result[0]['issues'],['TRÙNG KHÓA']);self.assertIsNone(result[0]['document'])
    def test_money_rate_date_discrepancies(self):
        source=dict(self.doc,currency='USD',rate='25000')
        doc=dict(source,total='1080000.01',rate='25100',date='2026-10-04')
        result,_=engine.reconcile([source],[doc]);self.assertIn('CHÊNH LỆCH TIỀN · total',result[0]['issues']);self.assertIn('CHÊNH TỶ GIÁ',result[0]['issues']);self.assertIn('CHÊNH NGÀY HÓA ĐƠN',result[0]['issues'])
        result,_=engine.reconcile([dict(source,rate='')],[doc]);self.assertIn('THIẾU TỶ GIÁ',result[0]['issues'])
    def test_untrusted_xml_and_invalid_money(self):
        with self.assertRaises(ValueError):engine.parse_xml(b'<!DOCTYPE x [<!ENTITY a "x">]><x>&a;</x>')
        with self.assertRaises(ValueError):engine.money('NaN')
        with self.assertRaises(ValueError):engine.parse_xml(b'<HDon/>')
    def test_csv_validation(self):
        records=engine.parse_csv(csv_bytes([self.doc]));self.assertEqual(records[0]['number'],'00000057')
        with self.assertRaises(ValueError):engine.parse_csv(csv_bytes([self.doc,self.doc]))
        with self.assertRaises(ValueError):engine.parse_csv(b'number,total\n57,123')
    def test_excel_roundtrip_six_sheets_and_sales(self):
        doc=dict(self.doc,kind='sales',net='',vat='',seller='=1+1')
        results,_=engine.reconcile([doc],[doc])
        wb=load_workbook(io.BytesIO(engine.excel({'name':'Test','tax':TAX},PERIOD,results,[],[])))
        self.assertEqual(len(wb.sheetnames),6);ws=wb['BangKe_MuaVao']
        self.assertEqual(ws['C2'].value,'00000057');self.assertEqual(ws['C2'].data_type,'s');self.assertEqual(ws['C2'].number_format,'@')
        self.assertEqual(ws['H2'].value,1080000);self.assertIsNone(ws['I2'].value)
        self.assertEqual(ws['E2'].value,'=1+1');self.assertEqual(ws['E2'].data_type,'s')
        self.assertEqual(ws.freeze_panes,'A2');self.assertEqual(ws.max_row,2)
        self.assertEqual(wb['ChiTiet_HoaDon']['C2'].value,self.doc['items'][0]['name'])
        self.assertEqual(ws['H2'].number_format,'#,##0')

class IntegrationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp=tempfile.TemporaryDirectory();cls.old=app.DATA;app.DATA=Path(cls.temp.name)
        cls.server=ThreadingHTTPServer(('127.0.0.1',0),app.Handler)
        cls.thread=threading.Thread(target=cls.server.serve_forever,daemon=True);cls.thread.start()
        cls.base=f'http://127.0.0.1:{cls.server.server_port}'
    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown();cls.server.server_close();cls.thread.join();app.DATA=cls.old;cls.temp.cleanup()
    def request(self,path,data=None,headers=None):
        request=Request(self.base+path,data=json.dumps(data).encode() if data is not None else None,headers=headers or {'Content-Type':'application/json','X-Invoice-App':'1'})
        with urlopen(request) as response:return response.read(),response.headers
    def test_complete_workflow_and_persistence(self):
        self.request('/api/companies',{'tax':TAX,'name':'Công ty kiểm thử'})
        self.request('/api/import',upload('cqt',csv_bytes([engine.parse_xml(XML)]),'cqt.csv'))
        data,_=self.request('/api/import',upload('xml',XML,'invoice.xml'));self.assertEqual(json.loads(data)['added'],1)
        data,_=self.request('/api/import',upload('xml',XML,'invoice-again.xml'));self.assertEqual(json.loads(data)['cached'],1)
        data,_=self.request(f'/api/report?tax={TAX}&period={PERIOD}');self.assertEqual(json.loads(data)['stats']['matched'],1)
        data,headers=self.request(f'/api/export?tax={TAX}&period={PERIOD}');self.assertIn('spreadsheetml',headers['Content-Type']);self.assertEqual(load_workbook(io.BytesIO(data))['BangKe_MuaVao']['C2'].value,'00000057')
        with app.connect() as db:self.assertEqual(app.report(db,TAX,PERIOD)['stats']['files'],1)
        # An invalid XML is tracked separately, without claiming an identified missing invoice.
        self.request('/api/import',upload('xml',b'<broken>','broken.xml'))
        data,_=self.request(f'/api/report?tax={TAX}&period={PERIOD}');report=json.loads(data);self.assertEqual(report['stats']['errors'],1);self.assertEqual(report['stats']['missing'],0)
        # Cross-company import is refused and existing records remain unchanged.
        with self.assertRaises(HTTPError) as exc:self.request('/api/import',upload('xml',XML.replace(TAX.encode(),b'9999999999'),'wrong.xml'))
        self.assertEqual(exc.exception.code,400)
        data,_=self.request(f'/api/report?tax={TAX}&period={PERIOD}');self.assertEqual(json.loads(data)['stats']['files'],1)
        data,_=self.request(f'/api/report?tax={TAX}&period=2026-09');self.assertEqual(json.loads(data)['stats']['total'],0)
    def test_origin_and_custom_header_required(self):
        with self.assertRaises(HTTPError):self.request('/api/demo',{'period':PERIOD},headers={'Content-Type':'application/json'})
        with self.assertRaises(HTTPError):self.request('/api/demo',{'period':PERIOD},headers={'Content-Type':'application/json','X-Invoice-App':'1','Origin':'https://example.com'})
    def test_demo_repeatable_and_static_available(self):
        self.request('/api/demo',{'period':PERIOD});self.request('/api/demo',{'period':PERIOD})
        data,_=self.request(f'/api/report?tax=0000000000&period={PERIOD}');stats=json.loads(data)['stats']
        self.assertEqual((stats['total'],stats['matched'],stats['exception'],stats['missing']),(10,7,3,2))
        html,headers=self.request('/');self.assertIn('Sổ Mua'.encode(),html);self.assertIn("frame-ancestors 'none'",headers['Content-Security-Policy'])

if __name__=='__main__':unittest.main()
