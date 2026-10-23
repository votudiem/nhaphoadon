import copy
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from openpyxl import Workbook,load_workbook
from pypdf import PdfReader
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import app
import documents
import engine
import workspace
from test_workflow import XML,TAX,PERIOD,csv_bytes,upload


def text_pdf():
    if 'TestingFont' not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont('TestingFont',str(documents.ROOT/'assets/DejaVuSans.ttf')))
    out=io.BytesIO();pdf=canvas.Canvas(out);pdf.setFont('TestingFont',11)
    lines=['HÓA ĐƠN GIÁ TRỊ GIA TĂNG','Ký hiệu: K26THT','Số hóa đơn: 00000057','Ngày: 03/10/2026','Đơn vị bán hàng: Công ty An Phát','Mã số thuế: 0311111111','Người mua: Công ty thử nghiệm','Mã số thuế: 0317557745','Cộng tiền hàng: 1.000.000','Tiền thuế GTGT: 80.000','Tổng tiền thanh toán: 1.080.000','Đơn vị tiền tệ: VND','Thiết bị giám sát hành trình GT-S7']
    for index,line in enumerate(lines):pdf.drawString(40,800-index*20,line)
    pdf.save();return out.getvalue()

class ExtendedTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.previous=app.DATA;app.DATA=Path(self.temp.name)/'data';self.db=app.connect()
        self.db.execute('INSERT INTO companies VALUES(?,?,0)',(TAX,'Công ty kiểm thử'))
        self.db.commit();self.doc=engine.parse_xml(XML)
    def tearDown(self):self.db.close();app.DATA=self.previous;self.temp.cleanup()
    def load_cqt(self):app.import_files(self.db,upload('cqt',csv_bytes([self.doc]),'cqt.csv'))
    def test_excel_cqt_mapping_missing_buyer_and_leading_zero(self):
        wb=Workbook();ws=wb.active;ws.append(['Danh sách hóa đơn']);ws.append(['MST người bán','Tên người bán','Ký hiệu hóa đơn','Số hóa đơn','Ngày lập','Tổng tiền chưa thuế','Tổng tiền thuế','Tổng tiền thanh toán']);ws.append(['0311111111','An Phát','K26THT',57,'03/10/2026',1000000,80000,1080000]);ws['D3'].number_format='00000000'
        out=io.BytesIO();wb.save(out);raw=out.getvalue()
        preview=documents.preview_table(raw,'cqt.xlsx');self.assertEqual(preview['header_row'],1)
        with self.assertRaises(ValueError):documents.parse_table(raw,'cqt.xlsx',TAX,PERIOD)
        records,filtered=documents.parse_table(raw,'cqt.xlsx',TAX,PERIOD,confirm_buyer=True)
        self.assertTrue(records[0]['buyer_asserted']);self.assertEqual(records[0]['number'],'00000057');self.assertEqual(records[0]['total'],'1080000');self.assertEqual(filtered,0)
    def test_pdf_text_requires_review_and_xml_preferred(self):
        self.load_cqt();raw=text_pdf();candidate=documents.parse_document(raw,'test.pdf')
        self.assertEqual(candidate['number'],'00000057');self.assertEqual(candidate['buyer_tax'],TAX);self.assertEqual(candidate['total'],'1080000');self.assertTrue(candidate['needs_review'])
        state=workspace.insert_document(self.db,TAX,PERIOD,raw,'test.pdf');self.assertEqual(state,'added')
        result=app.report(self.db,TAX,PERIOD);self.assertIn('CHỜ XÁC NHẬN TRÍCH XUẤT',result['results'][0]['issues'])
        workspace.insert_document(self.db,TAX,PERIOD,XML,'invoice.xml')
        result=app.report(self.db,TAX,PERIOD);self.assertEqual(result['stats']['matched'],1);self.assertEqual(len(result['results'][0]['document']['related_files']),2);self.assertEqual(result['stats']['orphans'],0)
    def test_manual_review_audited_and_reconciles(self):
        self.load_cqt();workspace.insert_document(self.db,TAX,PERIOD,text_pdf(),'invoice.pdf')
        row=self.db.execute('SELECT fingerprint FROM raw_files').fetchone();record=dict(self.doc)
        workspace.review(self.db,TAX,PERIOD,row['fingerprint'],record)
        self.assertEqual(app.report(self.db,TAX,PERIOD)['stats']['matched'],1)
        audit=self.db.execute("SELECT detail FROM audit WHERE event='XÁC NHẬN TRÍCH XUẤT'").fetchone();self.assertIn('before',json.loads(audit['detail']))
    def test_xml_to_pdf_unicode_and_derived_pdf_excluded(self):
        pdf=documents.pdf_representation(self.doc);reader=PdfReader(io.BytesIO(pdf));text=''.join(p.extract_text() for p in reader.pages)
        self.assertIn('00000057',text);self.assertIn('Thiết bị giám sát hành trình',text);self.assertIn('Bản thể hiện',text)
        self.assertEqual(workspace.insert_document(self.db,TAX,PERIOD,pdf,'representation.pdf'),'filtered')
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM documents').fetchone()[0],0)
    def test_retry_only_failed_files_and_keep_source(self):
        workspace.insert_document(self.db,TAX,PERIOD,XML,'valid.xml');workspace.insert_document(self.db,TAX,PERIOD,b'bad','bad.xml')
        with patch('workspace.documents.parse_document',side_effect=ValueError('still broken')) as parser:
            counts=workspace.retry_errors(self.db,TAX,PERIOD)
        self.assertEqual(parser.call_count,1);self.assertEqual(counts['failed'],1);self.assertEqual(self.db.execute("SELECT COUNT(*) FROM documents WHERE source='xml'").fetchone()[0],1)
    def test_period_close_locks_and_reopen_preserves_revision(self):
        self.load_cqt();workspace.insert_document(self.db,TAX,PERIOD,XML,'invoice.xml')
        closed=workspace.save_revision(self.db,TAX,PERIOD,app.report(self.db,TAX,PERIOD),True)
        self.assertTrue(closed['closed'])
        with self.assertRaises(ValueError):app.import_files(self.db,upload('xml',XML,'invoice.xml'))
        self.db.execute('UPDATE period_state SET closed=0 WHERE company=? AND period=?',(TAX,PERIOD))
        version=workspace.save_revision(self.db,TAX,PERIOD,app.report(self.db,TAX,PERIOD))
        self.assertEqual(version['version'],2);self.assertEqual(self.db.execute('SELECT COUNT(*) FROM revisions').fetchone()[0],2)
    def test_unresolved_exceptions_prevent_close(self):
        self.load_cqt()
        with self.assertRaises(ValueError):workspace.save_revision(self.db,TAX,PERIOD,app.report(self.db,TAX,PERIOD),True)
    def test_scan_filters_period_excludes_pdf_output_and_is_repeatable(self):
        folder=Path(self.temp.name)/'invoices';folder.mkdir();(folder/'a.xml').write_bytes(XML)
        (folder/'other-month.xml').write_bytes(XML.replace(b'2026-10-03',b'2026-09-03'))
        (folder/'_PDF_CQT').mkdir();(folder/'_PDF_CQT'/'generated.xml').write_bytes(XML)
        workspace.save_settings(self.db,TAX,{'source_folder':str(folder)})
        result=workspace.scan_folder(self.db,TAX,PERIOD);self.assertEqual((result['added'],result['filtered']),(1,1))
        result=workspace.scan_folder(self.db,TAX,PERIOD);self.assertEqual(result['cached'],1)
    def test_company_folders_cannot_overlap(self):
        folder=Path(self.temp.name)/'source';folder.mkdir();child=folder/'child';child.mkdir()
        workspace.save_settings(self.db,TAX,{'source_folder':str(folder)})
        self.db.execute('INSERT INTO companies VALUES(?,?,0)',('9999999999','Other'))
        with self.assertRaises(ValueError):workspace.save_settings(self.db,'9999999999',{'source_folder':str(child)})
    def test_approved_rules_only_modify_display_not_details(self):
        self.load_cqt();workspace.insert_document(self.db,TAX,PERIOD,XML,'invoice.xml')
        with self.assertRaises(ValueError):workspace.save_settings(self.db,TAX,{'rules':[{'keyword':'thiết bị','label':'Chi phí thiết bị','approved':False}]})
        workspace.save_settings(self.db,TAX,{'rules':[{'keyword':'thiết bị','label':'Chi phí thiết bị','approved':True}]})
        result=app.report(self.db,TAX,PERIOD);content=engine.excel(result['company'],PERIOD,result['results'],[],[]);wb=load_workbook(io.BytesIO(content));self.assertEqual(wb['BangKe_MuaVao']['G2'].value,'Chi phí thiết bị');self.assertEqual(wb['ChiTiet_HoaDon']['C2'].value,self.doc['items'][0]['name'])
    def test_changed_file_replaces_active_parse_but_keeps_original_archive(self):
        folder=Path(self.temp.name)/'files';folder.mkdir();file=folder/'invoice.xml';file.write_bytes(XML)
        workspace.save_settings(self.db,TAX,{'source_folder':str(folder)})
        workspace.scan_folder(self.db,TAX,PERIOD)
        file.write_bytes(XML.replace(b'1080000',b'1080001'))
        workspace.scan_folder(self.db,TAX,PERIOD)
        self.assertEqual(self.db.execute("SELECT COUNT(*) FROM documents WHERE source='xml'").fetchone()[0],1)
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM raw_files').fetchone()[0],2)
        data=json.loads(self.db.execute("SELECT body FROM documents WHERE source='xml'").fetchone()[0]);self.assertEqual(data['total'],'1080001')
    def test_batch_continues_after_failed_company_and_source_download_is_scoped(self):
        import base64,time,threading
        from http.server import ThreadingHTTPServer
        from urllib.request import Request,urlopen
        from urllib.error import HTTPError
        other='9999999999'
        broken=Path(self.temp.name)/'disconnected';broken.mkdir()
        good=Path(self.temp.name)/'other';good.mkdir();(good/'invoice.xml').write_bytes(XML.replace(TAX.encode(),other.encode()))
        self.db.execute('INSERT INTO companies VALUES(?,?,0)',(other,'Other'))
        workspace.save_settings(self.db,TAX,{'source_folder':str(broken)})
        workspace.save_settings(self.db,other,{'source_folder':str(good)})
        self.db.commit();broken.rmdir()
        server=ThreadingHTTPServer(('127.0.0.1',0),app.Handler);thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start();base=f'http://127.0.0.1:{server.server_port}'
        def call(path,data=None):
            request=Request(base+path,data=json.dumps(data).encode() if data is not None else None,headers={'Content-Type':'application/json','X-Invoice-App':'1'})
            with urlopen(request) as response:return response.read()
        try:
            with patch('routes.logging.exception'):
                identifier=json.loads(call('/api/batch',{'period':PERIOD}))['id']
                for _ in range(100):
                    job=json.loads(call('/api/job?id='+identifier))
                    if job['state']=='done':break
                    time.sleep(.02)
            self.assertEqual(job['state'],'done');self.assertEqual(len(job['results']),2);self.assertIn('error',job['results'][0]);self.assertEqual(job['results'][1]['result']['added'],1)
            fingerprint=self.db.execute('SELECT fingerprint FROM raw_files WHERE company=?',(other,)).fetchone()[0]
            self.assertIn(other.encode(),call(f'/api/file?tax={other}&period={PERIOD}&fingerprint={fingerprint}'))
            with self.assertRaises(HTTPError):call(f'/api/file?tax={TAX}&period={PERIOD}&fingerprint={fingerprint}')
            backup=call('/api/backup');self.assertTrue(backup.startswith(b'SQLite format 3'))
        finally:server.shutdown();server.server_close();thread.join()
    def test_corrupt_spreadsheet_is_an_actionable_error(self):
        with self.assertRaises(ValueError):documents.preview_table(b'not an excel file','bad.xlsx')
    @unittest.skipUnless(documents.shutil.which('tesseract'),'Tesseract is optional')
    def test_actual_local_ocr_process(self):
        from PIL import Image,ImageDraw,ImageFont
        image=Image.new('RGB',(1600,500),'white');draw=ImageDraw.Draw(image);font=ImageFont.truetype(str(documents.ROOT/'assets/DejaVuSans.ttf'),42)
        draw.text((40,60),'Invoice No: 00000057',font=font,fill='black');draw.text((40,140),'Tax code: 0311111111',font=font,fill='black');draw.text((40,220),'Total payment: 1,080,000 VND',font=font,fill='black')
        text,lang=documents.ocr_image(image);self.assertIn('00000057',text);self.assertIn('0311111111',text);self.assertIn('eng',lang)
    def test_local_ocr_fallback_uses_no_external_service(self):
        with patch('documents.ocr_image',return_value=('Mã số thuế: 0311111111\nMã số thuế: 0317557745','vie+eng')) as ocr:
            from PIL import Image
            out=io.BytesIO();Image.new('RGB',(100,100),'white').save(out,format='PNG')
            record=documents.parse_document(out.getvalue(),'scan.png');self.assertTrue(record['needs_review']);self.assertEqual(record['method'],'ocr:vie+eng');ocr.assert_called_once()

if __name__=='__main__':unittest.main()
